#!/usr/bin/env python3
"""Bouwt de twee spreadsheets voor /tools/inventory-profit-spreadsheet.

  python3 scripts/bouw_voorraadsheet.py

Schrijft backend/downloads/voorraad-winst-spreadsheet-nl.xlsx en
inventory-profit-spreadsheet-en.xlsx. Het resultaat staat in git, de server
bouwt niets zelf. Formules staan er echt in (geen vaste getallen); de proef
tests/test_voorraadsheet.py rekent ze na.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

UIT = Path(__file__).resolve().parent.parent / "backend" / "downloads"
RIJEN = 500  # regels met formules

TEKST = {
    "nl": dict(
        bestand="voorraad-winst-spreadsheet-nl.xlsx", blad1="Voorraad", blad2="Jaaroverzicht",
        kop=["Artikel", "Inkoopdatum", "Inkoopprijs", "Kanaal", "Datum online", "Verkoopdatum", "Verkoopprijs",
             "Verzendkosten (door jou)", "Platformkosten", "Winst", "Dagen online", "Status"],
        kanalen=["Marktplaats", "2dehands", "Vinted", "eBay", "Shopify", "Overig"],
        status=("Verkocht", "Online", "Voorraad"),
        jaar="Jaar", j_kop=["Platform", "Verkopen", "Omzet", "Winst", "DAC7: wordt doorgegeven?"],
        ja="Ja: het platform geeft je gegevens door aan de Belastingdienst", nee="Nee", totaal="Totaal",
        noot="DAC7-grens (Belastingdienst, per platform en kalenderjaar): 30 verkopen of meer, of meer dan 2.000 euro omzet. Doorgegeven is niet hetzelfde als belasting betalen. Dit is geen belastingadvies.",
        uitleg="Vul per artikel een regel in. De gele kolommen zijn voor jou; Winst, Dagen online en Status rekenen zichzelf uit.",
    ),
    "en": dict(
        bestand="inventory-profit-spreadsheet-en.xlsx", blad1="Inventory", blad2="Yearly overview",
        kop=["Item", "Purchase date", "Purchase price", "Marketplace", "Date listed", "Sale date", "Sale price",
             "Shipping cost (paid by you)", "Platform fees", "Profit", "Days online", "Status"],
        kanalen=["Marktplaats", "2dehands", "Vinted", "eBay", "Shopify", "Other"],
        status=("Sold", "Online", "In stock"),
        jaar="Year", j_kop=["Platform", "Sales", "Revenue", "Profit", "DAC7: gets reported?"],
        ja="Yes: the platform reports your details to the tax authority", nee="No", totaal="Total",
        noot="DAC7 limit (Dutch Belastingdienst, per platform and calendar year): 30 sales or more, or revenue above 2,000 euro. Being reported is not the same as paying tax. This is not tax advice.",
        uitleg="Fill in one row per item. The yellow columns are for you; Profit, Days online and Status calculate themselves.",
    ),
}


def bouw(taal: str) -> Path:
    t = TEKST[taal]
    wb = Workbook()
    ws = wb.active
    ws.title = t["blad1"]
    vet = Font(bold=True, color="FFFFFF")
    blauw = PatternFill("solid", fgColor="2563EB")
    geel = PatternFill("solid", fgColor="FEF9C3")
    grijs = PatternFill("solid", fgColor="F1F5F9")
    ws["A1"] = t["uitleg"]
    ws["A1"].font = Font(italic=True, color="64748B")
    kop = 3
    for i, naam in enumerate(t["kop"], 1):
        c = ws.cell(row=kop, column=i, value=naam)
        c.font, c.fill, c.alignment = vet, blauw, Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[kop].height = 32
    eerste, laatste = kop + 1, kop + RIJEN
    verkocht, online, voorraad = t["status"]
    for r in range(eerste, laatste + 1):
        for col in (1, 2, 3, 4, 5, 6, 7, 8, 9):
            ws.cell(row=r, column=col).fill = geel
        for col in (10, 11, 12):
            ws.cell(row=r, column=col).fill = grijs
        ws.cell(row=r, column=10, value=f'=IF(G{r}="","",G{r}-C{r}-H{r}-I{r})')
        ws.cell(row=r, column=11, value=f'=IF(E{r}="","",IF(F{r}="",TODAY(),F{r})-E{r})')
        ws.cell(row=r, column=12, value=f'=IF(F{r}<>"","{verkocht}",IF(E{r}<>"","{online}",IF(A{r}<>"","{voorraad}","")))')
        for col in (2, 5, 6):
            ws.cell(row=r, column=col).number_format = "dd-mm-yyyy"
        for col in (3, 7, 8, 9, 10):
            ws.cell(row=r, column=col).number_format = '#,##0.00 "€"' if taal == "nl" else '"€"#,##0.00'
        ws.cell(row=r, column=11).number_format = "0"
    dv = DataValidation(type="list", formula1='"' + ",".join(t["kanalen"]) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D{eerste}:D{laatste}")
    for col, w in zip("ABCDEFGHIJKL", (34, 13, 13, 14, 13, 13, 13, 16, 14, 12, 12, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=eerste, column=2)

    w2 = wb.create_sheet(t["blad2"])
    w2["A1"], w2["B1"] = t["jaar"], 2026
    w2["A1"].font = Font(bold=True)
    w2["B1"].fill = geel
    for i, naam in enumerate(t["j_kop"], 1):
        c = w2.cell(row=3, column=i, value=naam)
        c.font, c.fill = vet, blauw
    inv = f"'{t['blad1']}'!"
    d, f, g, j = (f"{inv}${c}${eerste}:${c}${laatste}" for c in "DFGJ")
    van, tot = 'DATE($B$1,1,1)', 'DATE($B$1,12,31)'
    for n, kanaal in enumerate(t["kanalen"]):
        r = 4 + n
        w2.cell(row=r, column=1, value=kanaal)
        w2.cell(row=r, column=2, value=f'=COUNTIFS({d},A{r},{f},">="&{van},{f},"<="&{tot})')
        w2.cell(row=r, column=3, value=f'=SUMIFS({g},{d},A{r},{f},">="&{van},{f},"<="&{tot})')
        w2.cell(row=r, column=4, value=f'=SUMIFS({j},{d},A{r},{f},">="&{van},{f},"<="&{tot})')
        w2.cell(row=r, column=5, value=f'=IF(OR(B{r}>=30,C{r}>2000),"{t["ja"]}","{t["nee"]}")')
    r = 4 + len(t["kanalen"])
    w2.cell(row=r, column=1, value=t["totaal"]).font = Font(bold=True)
    for col, letter in ((2, "B"), (3, "C"), (4, "D")):
        w2.cell(row=r, column=col, value=f"=SUM({letter}4:{letter}{r - 1})").font = Font(bold=True)
    for rr in range(4, r + 1):
        for col in (3, 4):
            w2.cell(row=rr, column=col).number_format = '#,##0.00 "€"' if taal == "nl" else '"€"#,##0.00'
    w2.cell(row=r + 2, column=1, value=t["noot"]).font = Font(italic=True, color="64748B")
    for col, w in zip("ABCDE", (18, 11, 14, 14, 60)):
        w2.column_dimensions[col].width = w
    UIT.mkdir(parents=True, exist_ok=True)
    pad = UIT / t["bestand"]
    wb.save(pad)
    return pad


if __name__ == "__main__":
    for taal in TEKST:
        print(bouw(taal))
